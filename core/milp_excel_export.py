from __future__ import annotations

import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

from core.milp_optimizer import BOXES_PER_PALLET, clean_mto, company_master, sku_master, sku_cooccurrence, solve_day


def _write_df(ws, df):
    for j, c in enumerate(df.columns, 1):
        ws.cell(1, j, c)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 300) + 1)]
        width = min(max(max((len(v) for v in vals), default=10) + 2, 10), 35)
        ws.column_dimensions[get_column_letter(col)].width = width


def build_milp_workbook(mto: pd.DataFrame, process4: pd.DataFrame) -> bytes:
    """Run the final MTO MILP for every supplied date and return an Excel output pack."""
    d = clean_mto(mto)
    required = ["SLOT_ID", "X_M", "Y_M", "DISTANCE_FROM_DOOR_M"]
    if not set(required).issubset(process4.columns):
        raise ValueError(f"Process-4 table missing {sorted(set(required) - set(process4.columns))}")
    slots = process4[required].copy().sort_values("DISTANCE_FROM_DOOR_M").reset_index(drop=True)
    if len(slots) == 0:
        raise ValueError("Process-4 selected area contains no pallet positions.")

    cm = company_master(d)
    sm = sku_master(d)
    co = sku_cooccurrence(d, 150)

    results = []
    allocations = []
    held_rows = []
    location_rows = []
    order_rows = []

    for dt in sorted(d.DATE.dt.normalize().dropna().unique()):
        r = solve_day(d, slots, dt)
        results.append(r["summary"])
        if not r["allocation"].empty:
            allocations.append(r["allocation"])
        if not r["held"].empty:
            held_rows.append(r["held"])
        if not r["orders"].empty:
            o = r["orders"].copy(); o["Date"] = pd.Timestamp(dt).date().isoformat(); order_rows.append(o)

        full = slots.copy()
        if not r["location_summary"].empty:
            full = full.merge(r["location_summary"], on="SLOT_ID", how="left", suffixes=("", "_calc"))
        for c in ["BOXES_STORED", "ORDERS", "COMPANIES", "UTILISATION_PCT"]:
            if c not in full:
                full[c] = 0
        full[["BOXES_STORED", "ORDERS", "COMPANIES", "UTILISATION_PCT"]] = full[["BOXES_STORED", "ORDERS", "COMPANIES", "UTILISATION_PCT"]].fillna(0)
        full["Date"] = pd.Timestamp(dt).date().isoformat()
        full["REMAINING_BOX_CAPACITY"] = BOXES_PER_PALLET - full["BOXES_STORED"]
        location_rows.append(full)

    daily = pd.DataFrame(results).sort_values("Date").reset_index(drop=True)
    alloc = pd.concat(allocations, ignore_index=True) if allocations else pd.DataFrame()
    held = pd.concat(held_rows, ignore_index=True) if held_rows else pd.DataFrame()
    locations = pd.concat(location_rows, ignore_index=True)
    orders = pd.concat(order_rows, ignore_index=True) if order_rows else pd.DataFrame()

    # Independent validation checks for the workbook.
    validation = []
    for _, s in daily.iterrows():
        a = alloc[alloc["Date"] == s["Date"]] if not alloc.empty else pd.DataFrame()
        h = held[held["Date"] == s["Date"]] if not held.empty else pd.DataFrame()
        stored = int(a["BOXES_STORED"].sum()) if not a.empty else 0
        held_boxes = int(h["BOXES_HELD"].sum()) if not h.empty else 0
        max_slot = float(a.groupby("SLOT_ID")["BOXES_STORED"].sum().max()) if not a.empty else 0.0
        validation.append({
            "Date": s["Date"],
            "MTO_boxes": int(s["MTO_boxes"]),
            "Stored_plus_Held": stored + held_boxes,
            "Conservation_OK": stored + held_boxes == int(s["MTO_boxes"]),
            "Max_boxes_any_pallet": max_slot,
            "Capacity_OK": max_slot <= BOXES_PER_PALLET,
            "Process4_positions": len(slots),
            "Process4_positions_OK": len(slots) == 93,
            "Solver_Optimal": "Optimal" in str(s["Solver_status"]),
        })
    validation = pd.DataFrame(validation)

    wb = Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("00_ReadMe")
    notes = [
        ("Project", "DATA-DRIVEN WAREHOUSE OPTIMIZATION AND IMPACT ANALYSIS"),
        ("Output format", "Excel-only output pack. No management dashboard is part of this final methodology."),
        ("Physical lineage", "Uses the existing Process-4 selected-area pallet positions and exact distance-from-door values; Processes 1–4 are not rerun."),
        ("Pallet capacity", f"{len(slots)} Process-4 positions × {BOXES_PER_PALLET} boxes per position"),
        ("MTO scope", "MTO only. MTA is excluded from this storage optimization."),
        ("Priority rule", "Smallest company set reaching at least 80% cumulative historical MTO box volume is marked Priority; all other companies remain eligible."),
        ("MILP Stage 1", "Maximize priority-company boxes stored."),
        ("MILP Stage 2", "Fix Stage 1 optimum, then maximize total boxes stored."),
        ("MILP Stage 3", "Fix Stages 1–2, then minimize box-distance from the door."),
        ("SKU grouping", "Movement score uses frequency and volume. Co-occurrence is a planning signal based on SKUs appearing in the same MTO order; it is not a physical compatibility certificate."),
        ("Historical interpretation", "Daily pallet equivalents are flow requirements for each date, not simultaneous inventory unless an inventory snapshot exists."),
    ]
    for i, (a, b) in enumerate(notes, 1):
        readme.cell(i, 1, a).font = Font(bold=True)
        readme.cell(i, 1).fill = PatternFill("solid", fgColor="D9EAF7")
        readme.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 110

    def add(name, df):
        ws = wb.create_sheet(name)
        _write_df(ws, df)
        return ws

    add("01_Validation", validation)
    add("02_Process4_Locations", slots)
    add("03_Company_Priority", cm)
    add("04_SKU_Movement", sm)
    add("05_SKU_Cooccurrence", co)
    add("06_Daily_Summary", daily)
    add("07_Daily_Location_Utilization", locations.sort_values(["Date", "DISTANCE_FROM_DOOR_M", "SLOT_ID"]))
    add("08_Order_Allocation", orders.sort_values(["Date", "INVOICE_ID"]))
    add("09_Storage_Allocation", alloc.sort_values(["Date", "DISTANCE_FROM_DOOR_M", "INVOICE_ID"]) if not alloc.empty else alloc)
    add("10_Held_Orders", held.sort_values(["Date", "BOXES_HELD"], ascending=[True, False]) if not held.empty else held)
    source = d.copy(); source["DATE"] = source["DATE"].dt.strftime("%Y-%m-%d")
    add("11_MTO_Cleaned_Input", source)

    kpi = pd.DataFrame([
        ["Process-4 pallet positions", len(slots)],
        ["Boxes per pallet", BOXES_PER_PALLET],
        ["MTO transaction lines", len(d)],
        ["Historical MTO boxes", int(d.BOXES.sum())],
        ["MTO dates", d.DATE.nunique()],
        ["Companies", len(cm)],
        ["Priority companies", int(cm.PRIORITY.sum())],
        ["Priority historical box share %", float(cm.loc[cm.PRIORITY, "MTO_BOXES"].sum() / cm.MTO_BOXES.sum() * 100)],
        ["SKUs", len(sm)],
        ["Peak MTO boxes", int(daily.MTO_boxes.max())],
        ["Peak pallet equivalent", int(daily.Required_pallet_equivalent.max())],
        ["Peak boxes held", int(daily.Boxes_held.max())],
        ["Overflow days vs Process-4 capacity", int(daily.Overflow.sum())],
        ["All daily MILP runs optimal", bool(validation.Solver_Optimal.all())],
    ], columns=["Metric", "Value"])
    add("12_Key_KPIs", kpi)

    # Excel-native charts with their source tables on the same sheet.
    charts = wb.create_sheet("13_Charts")
    charts["A1"] = "MTO MILP Analysis Charts"
    charts["A1"].font = Font(size=16, bold=True)
    source_cols = ["Date", "MTO_boxes", "Boxes_stored", "Boxes_held", "Required_pallet_equivalent"]
    for j, c in enumerate(source_cols, 1): charts.cell(3, j, c)
    for i, row in enumerate(daily[source_cols].itertuples(index=False, name=None), 4):
        for j, v in enumerate(row, 1): charts.cell(i, j, v)
    cats = Reference(charts, min_col=1, min_row=4, max_row=3 + len(daily))
    line = LineChart(); line.title = "Daily MTO boxes vs stored boxes"; line.y_axis.title = "Boxes"; line.x_axis.title = "Date"
    line.add_data(Reference(charts, min_col=2, max_col=3, min_row=3, max_row=3 + len(daily)), titles_from_data=True); line.set_categories(cats); line.height = 8; line.width = 16
    charts.add_chart(line, "G3")
    bar = BarChart(); bar.title = "Daily boxes held outside Process-4 capacity"; bar.y_axis.title = "Boxes"; bar.x_axis.title = "Date"
    bar.add_data(Reference(charts, min_col=4, min_row=3, max_row=3 + len(daily)), titles_from_data=True); bar.set_categories(cats); bar.height = 8; bar.width = 16
    charts.add_chart(bar, "G20")

    topc = cm.head(10)[["COMPANY_NAME", "MTO_BOXES"]]
    for j, c in enumerate(topc.columns, 7): charts.cell(3, j, c)
    for i, row in enumerate(topc.itertuples(index=False, name=None), 4):
        charts.cell(i, 7, row[0]); charts.cell(i, 8, row[1])
    bc = BarChart(); bc.type = "bar"; bc.title = "Top 10 companies by historical MTO boxes"; bc.x_axis.title = "Boxes"; bc.y_axis.title = "Company"
    bc.add_data(Reference(charts, min_col=8, min_row=3, max_row=3 + len(topc)), titles_from_data=True); bc.set_categories(Reference(charts, min_col=7, min_row=4, max_row=3 + len(topc))); bc.height = 8; bc.width = 15
    charts.add_chart(bc, "G37")

    tops = sm.head(15)[["ITEM_SIZE", "MOVEMENT_SCORE"]]
    for j, c in enumerate(tops.columns, 10): charts.cell(3, j, c)
    for i, row in enumerate(tops.itertuples(index=False, name=None), 4):
        charts.cell(i, 10, row[0]); charts.cell(i, 11, float(row[1]))
    bs = BarChart(); bs.type = "bar"; bs.title = "Top 15 SKUs by movement score"; bs.x_axis.title = "Movement score"; bs.y_axis.title = "SKU"
    bs.add_data(Reference(charts, min_col=11, min_row=3, max_row=3 + len(tops)), titles_from_data=True); bs.set_categories(Reference(charts, min_col=10, min_row=4, max_row=3 + len(tops))); bs.height = 9; bs.width = 15
    charts.add_chart(bs, "G54")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        if ws.max_row > 1:
            ws.row_dimensions[1].height = 24

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
