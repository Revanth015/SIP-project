"""Export and reporting engine for the Universal Warehouse Digital Twin."""

from pathlib import Path
import json
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment


def build_kpi_table(capacity_summary, demand_metrics, simulation_summary, layout_metrics=None):
    """Create one dashboard-friendly KPI table from the independent engines."""
    rows = [
        ("Theoretical Capacity", capacity_summary.get("theoretical_capacity"), "pallet slots"),
        ("Realistic Operational Capacity", capacity_summary.get("realistic_capacity"), "pallet slots"),
        ("Planning Occupancy", capacity_summary.get("planning_occupancy_pct"), "%"),
        ("Planning Capacity", capacity_summary.get("planning_capacity"), "pallet slots"),
        ("Average Daily Demand", demand_metrics.get("average_daily_demand"), "pallets/day"),
        ("Peak Daily Demand", demand_metrics.get("peak_demand"), "pallets/day"),
        ("95th Percentile Demand", demand_metrics.get("p95_demand"), "pallets/day"),
        ("Average Storage Usage", simulation_summary.get("average_storage_usage_pct", None), "%"),
        ("Average Distance", simulation_summary.get("average_distance_m"), "m/day"),
        ("Average Time", simulation_summary.get("average_time_min"), "min/day"),
        ("Average Fatigue Proxy", simulation_summary.get("average_fatigue"), "score"),
        ("Overflow Days", simulation_summary.get("overflow_days"), "days"),
    ]
    if layout_metrics:
        rows.extend([
            ("Average Slot Distance", layout_metrics.get("avg_distance_m"), "m"),
            ("Generated Slot Count", layout_metrics.get("capacity"), "slots"),
            ("Main Aisle", layout_metrics.get("main_aisle_m"), "m"),
            ("Cross Aisle", layout_metrics.get("cross_aisle_m"), "m"),
            ("Wall Clearance", layout_metrics.get("wall_clearance_m"), "m"),
            ("Pallet Orientation", layout_metrics.get("orientation"), "degrees"),
        ])
    return pd.DataFrame(rows, columns=["METRIC", "VALUE", "UNIT"])


def _write_df(writer, name, df):
    if df is None:
        return
    if isinstance(df, pd.Series):
        df = df.to_frame()
    elif not isinstance(df, pd.DataFrame):
        df = pd.DataFrame(df)
    df.to_excel(writer, sheet_name=name[:31], index=False)


def export_excel(
    output_path,
    kpi_table,
    capacity_table=None,
    daily_simulation=None,
    occupancy=None,
    slot_master=None,
    layout_search=None,
    assumptions=None,
    demand_metrics=None,
):
    """Export standard Digital Twin tables to one workbook."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_df(writer, "KPI_Dashboard", kpi_table)
        _write_df(writer, "Capacity", capacity_table)
        _write_df(writer, "Daily_Simulation", daily_simulation)
        _write_df(writer, "Slot_Occupancy", occupancy)
        _write_df(writer, "Slot_Master", slot_master)
        _write_df(writer, "Layout_Search", layout_search)
        _write_df(writer, "Demand_Metrics", pd.DataFrame([demand_metrics]) if demand_metrics is not None else None)
        _write_df(writer, "Assumptions", pd.DataFrame([(str(k), str(v)) for k, v in assumptions.items()], columns=["ASSUMPTION", "VALUE"]) if assumptions is not None else None)
    return output_path


def build_master_workbook_bytes(
    sheets: dict[str, pd.DataFrame],
    strategy_comparison: pd.DataFrame | None = None,
    daily_flow: pd.DataFrame | None = None,
    model1_summary: pd.DataFrame | None = None,
    title: str = "WAREHOUSEX — Master Digital Twin Output",
) -> bytes:
    """Build the complete project output workbook in memory, including charts.

    The caller supplies all model outputs, so the application itself—not a
    separate post-processing script—generates the final deliverable.
    """
    out = BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        used = set()
        for raw_name, df in sheets.items():
            if df is None:
                continue
            name = str(raw_name)[:31] or "Output"
            base = name
            i = 1
            while name in used:
                suffix = f"_{i}"
                name = base[: 31 - len(suffix)] + suffix
                i += 1
            used.add(name)
            if isinstance(df, pd.Series):
                df = df.to_frame()
            elif not isinstance(df, pd.DataFrame):
                df = pd.DataFrame(df)
            df.to_excel(writer, sheet_name=name, index=False)

    wb = load_workbook(BytesIO(out.getvalue()))
    dashboard = wb.create_sheet("Dashboard", 0)
    dashboard["A1"] = title
    dashboard["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    dashboard.merge_cells("A1:H1")
    dashboard["A3"] = "Generated by the application"
    dashboard["A3"].font = Font(bold=True)
    dashboard["B3"] = "All tables and charts in this workbook are generated from the current session outputs."
    dashboard["B3"].alignment = Alignment(wrap_text=True)

    row = 5
    if model1_summary is not None and not model1_summary.empty:
        dashboard.cell(row, 1, "Model 1 headline results")
        dashboard.cell(row, 1).font = Font(bold=True, size=13)
        for j, col in enumerate(model1_summary.columns, 1):
            dashboard.cell(row + 1, j, col).font = Font(bold=True)
        for i, vals in enumerate(model1_summary.itertuples(index=False), row + 2):
            for j, val in enumerate(vals, 1):
                dashboard.cell(i, j, val)
        row += len(model1_summary) + 4

    if strategy_comparison is not None and not strategy_comparison.empty:
        start = row
        dashboard.cell(start, 1, "Model 2 strategy comparison")
        dashboard.cell(start, 1).font = Font(bold=True, size=13)
        for j, col in enumerate(strategy_comparison.columns, 1):
            dashboard.cell(start + 1, j, col).font = Font(bold=True)
        for i, vals in enumerate(strategy_comparison.itertuples(index=False), start + 2):
            for j, val in enumerate(vals, 1):
                dashboard.cell(i, j, val)

        # Line coverage chart
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Historical MTO line coverage by strategy"
        chart.x_axis.title = "Line coverage (%)"
        chart.y_axis.title = "Strategy"
        cols = list(strategy_comparison.columns)
        if "Strategy" in cols and "Line coverage %" in cols:
            sx = cols.index("Strategy") + 1
            sy = cols.index("Line coverage %") + 1
            data = Reference(dashboard, min_col=sy, min_row=start + 1, max_row=start + 1 + len(strategy_comparison))
            cats = Reference(dashboard, min_col=sx, min_row=start + 2, max_row=start + 1 + len(strategy_comparison))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 7
            chart.width = 12
            dashboard.add_chart(chart, f"H{start}")

        # Belt coverage chart
        chart2 = BarChart()
        chart2.type = "bar"
        chart2.title = "Historical MTO belt coverage by strategy"
        chart2.x_axis.title = "Belt coverage (%)"
        if "Strategy" in cols and "Belt coverage %" in cols:
            sx = cols.index("Strategy") + 1
            sy = cols.index("Belt coverage %") + 1
            data = Reference(dashboard, min_col=sy, min_row=start + 1, max_row=start + 1 + len(strategy_comparison))
            cats = Reference(dashboard, min_col=sx, min_row=start + 2, max_row=start + 1 + len(strategy_comparison))
            chart2.add_data(data, titles_from_data=True)
            chart2.set_categories(cats)
            chart2.height = 7
            chart2.width = 12
            dashboard.add_chart(chart2, f"H{start + 15}")

        # Balanced score chart
        chart3 = BarChart()
        chart3.type = "bar"
        chart3.title = "Balanced performance score"
        chart3.x_axis.title = "Score"
        if "Strategy" in cols and "Balanced_Performance_Score" in cols:
            sx = cols.index("Strategy") + 1
            sy = cols.index("Balanced_Performance_Score") + 1
            data = Reference(dashboard, min_col=sy, min_row=start + 1, max_row=start + 1 + len(strategy_comparison))
            cats = Reference(dashboard, min_col=sx, min_row=start + 2, max_row=start + 1 + len(strategy_comparison))
            chart3.add_data(data, titles_from_data=True)
            chart3.set_categories(cats)
            chart3.height = 7
            chart3.width = 12
            dashboard.add_chart(chart3, f"A{start + 15}")

    if daily_flow is not None and not daily_flow.empty and "DATE" in daily_flow.columns:
        daily_start = max(row + 25, 50)
        dashboard.cell(daily_start, 1, "Daily MTO flow")
        dashboard.cell(daily_start, 1).font = Font(bold=True, size=13)
        for j, col in enumerate(daily_flow.columns, 1):
            dashboard.cell(daily_start + 1, j, col).font = Font(bold=True)
        for i, vals in enumerate(daily_flow.itertuples(index=False), daily_start + 2):
            for j, val in enumerate(vals, 1):
                dashboard.cell(i, j, val)
        cols = list(daily_flow.columns)
        if "MTO_BELTS" in cols:
            chart4 = LineChart()
            chart4.title = "Daily MTO belts"
            chart4.y_axis.title = "MTO belts"
            chart4.x_axis.title = "Date"
            sy = cols.index("MTO_BELTS") + 1
            sx = cols.index("DATE") + 1
            data = Reference(dashboard, min_col=sy, min_row=daily_start + 1, max_row=daily_start + 1 + len(daily_flow))
            cats = Reference(dashboard, min_col=sx, min_row=daily_start + 2, max_row=daily_start + 1 + len(daily_flow))
            chart4.add_data(data, titles_from_data=True)
            chart4.set_categories(cats)
            chart4.height = 8
            chart4.width = 16
            dashboard.add_chart(chart4, f"A{daily_start + 2}")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            max_len = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells) + 2, 45)
            ws.column_dimensions[column_cells[0].column_letter].width = max_len

    final = BytesIO()
    wb.save(final)
    return final.getvalue()


def export_summary_json(output_path, summary):
    """Export a machine-readable summary for integrations or future APIs."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, default=str)
    return output_path
